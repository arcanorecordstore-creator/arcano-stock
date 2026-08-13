# -*- coding: utf-8 -*-
"""Informe COMPLETO de contabilidad en Excel (multi-hoja, con graficos).

Refleja toda la pestana de /admin/contabilidad: resumen del periodo, ventas,
egresos, pedidos, evolucion de 6 meses (caja), Arcano vs coleccion personal,
top sellos y discos, stock frenado y dinero por cobrar. Devuelve un io.BytesIO.
"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

import metrics
from models import Disc, Pedido, MovimientoCaja, Setting

# --- Paleta / estilos ---
NEGRO = "0A0A0A"
VERDE = "0F7A53"
ROJO = "C0392B"
GRIS = "6B6B6B"
BEIGE = "F2F1EC"
BLANCO = "FFFFFF"

F_TITULO = Font(name="Calibri", size=16, bold=True, color=NEGRO)
F_SUB = Font(name="Calibri", size=9, color=GRIS)
F_H2 = Font(name="Calibri", size=12, bold=True, color=NEGRO)
F_TH = Font(name="Calibri", size=10, bold=True, color=BLANCO)
F_CELL = Font(name="Calibri", size=10, color="1A1A1A")
F_BOLD = Font(name="Calibri", size=10, bold=True, color=NEGRO)
FILL_TH = PatternFill("solid", fgColor=NEGRO)
FILL_BEIGE = PatternFill("solid", fgColor=BEIGE)
FILL_ZEBRA = PatternFill("solid", fgColor="FAF9F6")
RIGHT = Alignment(horizontal="right")
CENTER = Alignment(horizontal="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="E3E2DC")
BORDER = Border(bottom=THIN)


def _money(n):
    try:
        return float(n or 0)
    except (TypeError, ValueError):
        return 0.0


def _per_label(periodo):
    return {
        "mes_actual": "Este mes", "mes_pasado": "Mes pasado", "anio": "Este ano",
        "todo": "Todo el historial", "custom": "Periodo elegido",
    }.get(periodo, periodo)


def _titulo_hoja(ws, texto, sub=None):
    ws["A1"] = texto
    ws["A1"].font = F_TITULO
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_SUB
    ws.sheet_view.showGridLines = False


def _encabezados(ws, fila, headers, anchos=None):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=fila, column=i, value=h)
        c.font = F_TH
        c.fill = FILL_TH
        c.alignment = LEFT if i == 1 else CENTER
        if anchos and i <= len(anchos):
            ws.column_dimensions[get_column_letter(i)].width = anchos[i - 1]
    return fila + 1


def _fila_datos(ws, fila, valores, money_cols=(), zebra=False):
    for i, v in enumerate(valores, start=1):
        c = ws.cell(row=fila, column=i, value=v)
        c.font = F_CELL
        c.border = BORDER
        if zebra:
            c.fill = FILL_ZEBRA
        if i in money_cols:
            c.number_format = '#,##0'
            c.alignment = RIGHT
        elif i > 1 and isinstance(v, (int, float)):
            c.alignment = CENTER
    return fila + 1


def generar_excel(periodo="mes_actual", desde=None, hasta=None):
    conta = metrics.contabilidad(periodo, desde, hasta)
    data = metrics.get_all_metrics()
    kpis = data["kpis"]
    por_cobrar = metrics.dinero_por_cobrar()
    seis_m = metrics.ingresos_vs_egresos_6m()
    arc_per = metrics.arcano_vs_personal_6m()
    recup = metrics.recuperacion_inversion()
    frenado = metrics.stock_frenado()

    inicio, fin = metrics._rango_periodo(periodo, desde, hasta)
    inicio_dt = datetime(inicio.year, inicio.month, inicio.day)
    fin_dt = datetime(fin.year, fin.month, fin.day)

    # Caja actual (saldo)
    movs = MovimientoCaja.query.all()
    try:
        saldo_ini = float(Setting.get("caja_saldo_inicial") or 0)
    except (TypeError, ValueError):
        saldo_ini = 0.0
    ing_c = sum(m.monto_uyu or 0 for m in movs if m.tipo == "ingreso")
    egr_c = sum(m.monto_uyu or 0 for m in movs if m.tipo == "egreso")
    saldo = saldo_ini + ing_c - egr_c

    stock = (
        Disc.query
        .filter((Disc.estado == "Disponible") | ((Disc.estado == "Pre-order") & (Disc.pedido_id.isnot(None))))
        .filter((Disc.origen != "personal") | (Disc.origen.is_(None)))
        .all()
    )
    mercaderia = sum((d.costo_final_uyu or 0) * (d.stock or 1) for d in stock)

    wb = Workbook()

    # ==================== HOJA 1: RESUMEN ====================
    ws = wb.active
    ws.title = "Resumen"
    _titulo_hoja(
        ws, "ARCANO - Contabilidad",
        "%s (%s a %s) - generado el %s - en pesos uruguayos (UYU)"
        % (_per_label(periodo), conta["inicio"], conta["fin"], datetime.now().strftime("%d/%m/%Y")),
    )
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 20

    fila = 4
    ws.cell(row=fila, column=1, value="RESULTADO DEL PERIODO").font = F_H2
    fila += 1
    resumen_rows = [
        ("Ingresos (ventas)", conta["ingresos_uyu"]),
        ("- Costo de los vendidos (Arcano)", -conta["costo_discos_uyu"]),
        ("= Ganancia bruta", conta["ganancia_bruta_uyu"]),
        ("- Stock no vendido (Arcano)", -conta["costo_stock_no_vendido_uyu"]),
        ("- Egresos (gastos)", -conta["egresos_uyu"]),
        ("= GANANCIA REAL", conta["ganancia_real_uyu"]),
    ]
    for label, val in resumen_rows:
        c1 = ws.cell(row=fila, column=1, value=label)
        c2 = ws.cell(row=fila, column=2, value=_money(val))
        c2.number_format = '#,##0'
        c2.alignment = RIGHT
        destaca = label.startswith("=")
        c1.font = F_BOLD if destaca else F_CELL
        c2.font = F_BOLD if destaca else F_CELL
        if destaca:
            c1.fill = FILL_BEIGE
            c2.fill = FILL_BEIGE
        fila += 1

    comp = conta.get("comparacion")
    if comp and comp.get("var_ganancia_pct") is not None:
        signo = "subio" if comp["mejoro"] else "bajo"
        ws.cell(row=fila, column=1,
                value="Vs periodo anterior: %s %s%% (antes %s)"
                % (signo, abs(comp["var_ganancia_pct"]), int(round(_money(comp["ganancia_real_ant"]))))
                ).font = F_SUB
        fila += 1
    fila += 1

    ws.cell(row=fila, column=1, value="POR ORIGEN").font = F_H2
    fila += 1
    for label, val in [
        ("Ingresos de Arcano", conta["ingresos_arcano_uyu"]),
        ("Ingresos de tu coleccion personal", conta["ingresos_personal_uyu"]),
        ("Unidades vendidas (total)", conta["unidades_vendidas"]),
        ("Unidades de tu coleccion", conta["unidades_personal"]),
    ]:
        ws.cell(row=fila, column=1, value=label).font = F_CELL
        c2 = ws.cell(row=fila, column=2, value=_money(val) if "Ingresos" in label else val)
        if "Ingresos" in label:
            c2.number_format = '#,##0'
        c2.alignment = RIGHT
        c2.font = F_CELL
        fila += 1
    fila += 1

    ws.cell(row=fila, column=1, value="ESTADO ACTUAL DEL NEGOCIO (HOY)").font = F_H2
    fila += 1
    for label, val, fmt in [
        ("Plata en caja", saldo, 'money'),
        ("En mercaderia (stock a costo)", mercaderia, 'money'),
        ("Inversion en pedidos recuperada", recup["pct_real"], 'pct'),
        ("Invertido en pedidos", recup["invertido"], 'money'),
        ("Recuperado de pedidos", recup["recuperado"], 'money'),
        ("Falta recuperar", recup["falta"], 'money'),
    ]:
        ws.cell(row=fila, column=1, value=label).font = F_CELL
        c2 = ws.cell(row=fila, column=2)
        if fmt == 'pct':
            c2.value = "%s%%" % round(val)
        else:
            c2.value = _money(val)
            c2.number_format = '#,##0'
        c2.alignment = RIGHT
        c2.font = F_CELL
        fila += 1
    fila += 1

    ws.cell(row=fila, column=1, value="INDICADORES (TODO EL HISTORIAL)").font = F_H2
    fila += 1
    for label, val, fmt in [
        ("Stock disponible (titulos)", kpis["stock_disponible_titulos"], 'int'),
        ("Stock disponible (unidades)", kpis["stock_disponible_unidades"], 'int'),
        ("Inversion en stock", kpis["inversion_uyu"], 'money'),
        ("Valor de venta del stock", kpis["valor_venta_uyu"], 'money'),
        ("Ganancia potencial del stock", kpis["ganancia_potencial_uyu"], 'money'),
        ("Ventas historicas (unidades)", kpis["ventas_total_unidades"], 'int'),
        ("Facturado historico", kpis["facturado_uyu"], 'money'),
        ("Ganancia realizada historica", kpis["ganancia_realizada_uyu"], 'money'),
        ("Margen promedio", kpis["margen_promedio"], 'pct'),
        ("Por cobrar (reservados)", por_cobrar["total_uyu"], 'money'),
    ]:
        ws.cell(row=fila, column=1, value=label).font = F_CELL
        c2 = ws.cell(row=fila, column=2)
        if fmt == 'pct':
            c2.value = "%s%%" % val
        elif fmt == 'money':
            c2.value = _money(val)
            c2.number_format = '#,##0'
        else:
            c2.value = val
        c2.alignment = RIGHT
        c2.font = F_CELL
        fila += 1

    # ==================== HOJA 2: VENTAS ====================
    ws = wb.create_sheet("Ventas del periodo")
    _titulo_hoja(ws, "Ventas del periodo", "%s a %s" % (conta["inicio"], conta["fin"]))
    f = _encabezados(ws, 4,
                     ["Fecha", "Artista", "Titulo", "Sello", "Origen", "Precio", "Copias", "Costo", "Ganancia"],
                     anchos=[12, 26, 30, 20, 12, 12, 8, 12, 12])
    vendidos = (
        Disc.query.filter(Disc.estado == "Vendido")
        .filter(Disc.fecha_venta >= inicio_dt, Disc.fecha_venta < fin_dt)
        .order_by(Disc.fecha_venta.asc()).all()
    )
    tot_precio = tot_costo = tot_gan = 0
    for i, d in enumerate(vendidos):
        precio = (d.precio_mostrar or d.precio_uyu or 0)
        # Los discos personales no le costaron nada al negocio: costo 0 (igual que
        # en el Resumen, que excluye el costo de los personales de la ganancia).
        costo = 0 if (d.origen or "arcano") == "personal" else (d.costo_final_uyu or 0)
        gan = (precio - costo) * d.copias_facturadas
        tot_precio += precio * d.copias_facturadas
        tot_costo += costo * d.copias_facturadas
        tot_gan += gan
        f = _fila_datos(ws, f, [
            d.fecha_venta.strftime("%Y-%m-%d") if d.fecha_venta else "",
            d.artista or "", d.titulo or "", d.sello or "", (d.origen or "arcano"),
            round(precio), d.copias_facturadas, round(costo), round(gan),
        ], money_cols=(6, 8, 9), zebra=(i % 2 == 1))
    if vendidos:
        ws.cell(row=f, column=1, value="TOTAL").font = F_BOLD
        for col, val in [(6, tot_precio), (8, tot_costo), (9, tot_gan)]:
            c = ws.cell(row=f, column=col, value=round(val))
            c.font = F_BOLD
            c.number_format = '#,##0'
            c.alignment = RIGHT
    else:
        ws.cell(row=f, column=1, value="Sin ventas en el periodo.").font = F_SUB

    # ==================== HOJA 3: EGRESOS ====================
    from models import Egreso
    ws = wb.create_sheet("Egresos del periodo")
    _titulo_hoja(ws, "Egresos del periodo", "%s a %s" % (conta["inicio"], conta["fin"]))
    egresos = (
        Egreso.query.filter(Egreso.fecha >= inicio, Egreso.fecha < fin)
        .order_by(Egreso.fecha.asc()).all()
    )
    f = _encabezados(ws, 4, ["Fecha", "Categoria", "Descripcion", "Monto", "Moneda", "TC", "Monto UYU"],
                     anchos=[12, 18, 34, 12, 10, 10, 14])
    tot_egr = 0
    for i, e in enumerate(egresos):
        tot_egr += (e.monto_uyu or 0)
        f = _fila_datos(ws, f, [
            e.fecha.strftime("%Y-%m-%d"), e.categoria or "", e.descripcion or "",
            round(e.monto or 0, 2), e.moneda, round(e.tc_usado or 0, 2), round(e.monto_uyu or 0),
        ], money_cols=(7,), zebra=(i % 2 == 1))
    if egresos:
        ws.cell(row=f, column=1, value="TOTAL").font = F_BOLD
        c = ws.cell(row=f, column=7, value=round(tot_egr))
        c.font = F_BOLD
        c.number_format = '#,##0'
        c.alignment = RIGHT
        f += 2
    else:
        ws.cell(row=f, column=1, value="Sin egresos en el periodo.").font = F_SUB
        f += 2
    # Egresos por categoria
    if conta["egresos_por_categoria"]:
        ws.cell(row=f, column=1, value="POR CATEGORIA").font = F_H2
        f += 1
        f = _encabezados(ws, f, ["Categoria", "Total UYU"], anchos=[24, 16])
        for cat, monto in conta["egresos_por_categoria"]:
            f = _fila_datos(ws, f, [cat or "(sin categoria)", round(monto)], money_cols=(2,))

    # ==================== HOJA 4: PEDIDOS ====================
    ws = wb.create_sheet("Pedidos")
    _titulo_hoja(ws, "Pedidos (compras de stock)", "Estado a hoy")
    peds = Pedido.query.order_by(Pedido.fecha.desc()).all()
    f = _encabezados(ws, 4,
                     ["Numero", "Fecha", "Proveedor", "Costo", "Recaudado", "Resultado",
                      "Margen %", "Rotacion %", "Discos", "Vendidos"],
                     anchos=[16, 12, 14, 14, 14, 14, 10, 10, 9, 9])

    def _g(obj, attr, default=0):
        try:
            v = getattr(obj, attr)
            return v if v is not None else default
        except Exception:
            return default
    for i, p in enumerate(peds):
        f = _fila_datos(ws, f, [
            p.numero or "-", p.fecha.strftime("%Y-%m-%d") if p.fecha else "",
            p.proveedor or "-", round(_g(p, "costo_total_uyu")), round(_g(p, "recaudado_uyu")),
            round(_g(p, "resultado_uyu")), round(_g(p, "margen_pct"), 1), round(_g(p, "rotacion_pct"), 1),
            _g(p, "discos_asignados"), _g(p, "discos_vendidos"),
        ], money_cols=(4, 5, 6), zebra=(i % 2 == 1))
    if not peds:
        ws.cell(row=f, column=1, value="No hay pedidos cargados.").font = F_SUB

    # ==================== HOJA 5: 6 MESES (CAJA) ====================
    ws = wb.create_sheet("6 meses (caja)")
    _titulo_hoja(ws, "Plata que entro vs salio (6 meses)",
                 "Mirada de caja: los egresos incluyen los pedidos. Neto = entro - salio.")
    f0 = 4
    f = _encabezados(ws, f0, ["Mes", "Ingresos", "Egresos (con pedidos)", "Gastos operativos", "Pedidos", "Neto"],
                     anchos=[10, 14, 20, 18, 14, 14])
    primera = f
    for i, m in enumerate(seis_m):
        f = _fila_datos(ws, f, [
            m["label"], round(m["ingresos_uyu"]), round(m["egresos_uyu"]),
            round(m["egresos_opex_uyu"]), round(m["pedidos_uyu"]), round(m["neto_uyu"]),
        ], money_cols=(2, 3, 4, 5, 6), zebra=(i % 2 == 1))
    # Grafico de barras: ingresos vs egresos
    if seis_m:
        chart = BarChart()
        chart.type = "col"
        chart.title = "Ingresos vs Egresos por mes"
        chart.height = 8
        chart.width = 18
        cats = Reference(ws, min_col=1, min_row=primera, max_row=f - 1)
        dat_ing = Reference(ws, min_col=2, min_row=f0, max_row=f - 1)
        dat_egr = Reference(ws, min_col=3, min_row=f0, max_row=f - 1)
        chart.add_data(dat_ing, titles_from_data=True)
        chart.add_data(dat_egr, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, "H4")

    # ==================== HOJA 6: ARCANO VS PERSONAL ====================
    ws = wb.create_sheet("Arcano vs Personal")
    _titulo_hoja(ws, "Ventas: Arcano vs tu coleccion personal (6 meses)",
                 "Cuanto se sostiene el negocio por si solo.")
    f0 = 4
    f = _encabezados(ws, f0, ["Mes", "Arcano", "Personal"], anchos=[10, 14, 14])
    primera = f
    for i, m in enumerate(arc_per):
        f = _fila_datos(ws, f, [m["label"], round(m["arcano"]), round(m["personal"])],
                        money_cols=(2, 3), zebra=(i % 2 == 1))
    if arc_per:
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "stacked"
        chart.overlap = 100
        chart.title = "Arcano vs Personal por mes"
        chart.height = 8
        chart.width = 16
        cats = Reference(ws, min_col=1, min_row=primera, max_row=f - 1)
        dat = Reference(ws, min_col=2, max_col=3, min_row=f0, max_row=f - 1)
        chart.add_data(dat, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, "F4")

    # ==================== HOJA 7: TOP SELLOS Y DISCOS ====================
    ws = wb.create_sheet("Top sellos y discos")
    _titulo_hoja(ws, "Top sellos y discos")
    f = 4
    ws.cell(row=f, column=1, value="TOP SELLOS POR INVERSION (stock)").font = F_H2
    f += 1
    f = _encabezados(ws, f, ["Sello", "Titulos", "Unidades", "Inversion UYU"], anchos=[24, 10, 10, 16])
    for i, s in enumerate(data["top_sellos_inversion"]):
        f = _fila_datos(ws, f, [s["sello"], s["titulos"], s["unidades"], round(s["inversion_uyu"])],
                        money_cols=(4,), zebra=(i % 2 == 1))
    f += 1
    ws.cell(row=f, column=1, value="TOP SELLOS POR ROTACION (vendido)").font = F_H2
    f += 1
    f = _encabezados(ws, f, ["Sello", "Titulos vendidos", "Unidades", "Facturado UYU"], anchos=[24, 14, 10, 16])
    for i, s in enumerate(data["top_sellos_rotacion"]):
        f = _fila_datos(ws, f, [s["sello"], s["titulos_vendidos"], s["unidades_vendidas"], round(s["facturado_uyu"])],
                        money_cols=(4,), zebra=(i % 2 == 1))
    f += 1
    ws.cell(row=f, column=1, value="TOP DISCOS MAS RENTABLES (vendidos)").font = F_H2
    f += 1
    f = _encabezados(ws, f, ["Artista", "Titulo", "Sello", "Precio", "Ganancia", "Margen %"],
                     anchos=[24, 28, 18, 12, 12, 10])
    for i, d in enumerate(data["top_discos_rentables"]):
        f = _fila_datos(ws, f, [
            d["artista"] or "", d["titulo"] or "", d["sello"] or "",
            round(d["precio_uyu"] or 0), round(d["ganancia_uyu"] or 0), round(d["margen_pct"] or 0, 1),
        ], money_cols=(4, 5), zebra=(i % 2 == 1))

    # ==================== HOJA 8: STOCK FRENADO ====================
    ws = wb.create_sheet("Stock frenado")
    _titulo_hoja(ws, "Stock que no rota",
                 "%d disco(s) parados > %d dias, con %s quietos."
                 % (frenado["frenados_count"], frenado["umbral"], int(round(frenado["plata_frenada"]))))
    f = _encabezados(ws, 4, ["Dias en stock", "Artista", "Titulo", "Sello", "Costo", "Precio"],
                     anchos=[14, 24, 28, 18, 12, 12])
    for i, fila_f in enumerate(frenado["filas"]):
        d = fila_f["disc"]
        f = _fila_datos(ws, f, [
            fila_f["dias"] if fila_f["dias"] is not None else "-",
            d.artista or "", d.titulo or "", d.sello or "",
            round(fila_f["costo"]), round(fila_f["precio"]),
        ], money_cols=(5, 6), zebra=(i % 2 == 1))
    if not frenado["filas"]:
        ws.cell(row=f, column=1, value="No hay stock de Arcano disponible.").font = F_SUB

    # ==================== HOJA 9: POR COBRAR ====================
    ws = wb.create_sheet("Por cobrar")
    _titulo_hoja(ws, "Dinero por cobrar (reservados)",
                 "%d disco(s) reservados, total %s."
                 % (por_cobrar["titulos"], int(round(por_cobrar["total_uyu"]))))
    f = _encabezados(ws, 4, ["Artista", "Titulo", "Sello", "Precio"], anchos=[26, 30, 20, 14])
    for i, d in enumerate(por_cobrar["discos"]):
        f = _fila_datos(ws, f, [
            d.artista or "", d.titulo or "", d.sello or "",
            round(d.precio_mostrar or d.precio_uyu or 0),
        ], money_cols=(4,), zebra=(i % 2 == 1))
    if not por_cobrar["discos"]:
        ws.cell(row=f, column=1, value="No hay discos reservados.").font = F_SUB

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
